import io

from flask import session, send_file

from db import get_registrations, get_users, get_solar_requests, audit_log, get_db, dict_cur

from . import admin_bp, require_admin


def _xl_response(wb, filename):
    from openpyxl.styles import Font, PatternFill, Alignment
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return send_file(buf,
                     mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                     as_attachment=True,
                     download_name=filename)


def _xl_header(ws, headers):
    from openpyxl.styles import Font, PatternFill, Alignment
    fill = PatternFill("solid", fgColor="0D3B27")
    font = Font(color="FFFFFF", bold=True, size=11)
    for col, h in enumerate(headers, 1):
        c = ws.cell(row=1, column=col, value=h)
        c.fill = fill
        c.font = font
        c.alignment = Alignment(horizontal="center")
    ws.freeze_panes = "A2"


@admin_bp.route("/admin/export/registrations")
def export_registrations():
    resp = require_admin()
    if resp:
        return resp
    from openpyxl import Workbook
    regs = get_registrations()
    wb = Workbook(); ws = wb.active; ws.title = "Registrations"
    headers = ["Ref #", "Form Type", "Full Name", "Email", "Contact",
               "Company", "Training Type", "Status", "Submitted At"]
    _xl_header(ws, headers)
    for r in regs:
        ws.append([r.get("id"), r.get("form_type"), r.get("full_name"),
                   r.get("email"), r.get("contact_number"), r.get("company_name"),
                   r.get("training_type"), r.get("status"), str(r.get("submitted_at", ""))])
    audit_log(session.get("user_id"), "export_report", "report", None, "Downloaded registrations.xlsx")
    return _xl_response(wb, "registrations.xlsx")


@admin_bp.route("/admin/export/users")
def export_users():
    resp = require_admin()
    if resp:
        return resp
    from openpyxl import Workbook
    users = get_users()
    wb = Workbook(); ws = wb.active; ws.title = "Users"
    headers = ["ID", "Role", "First Name", "Last Name", "Email",
               "Contact", "Company", "Payment Status", "Created At"]
    _xl_header(ws, headers)
    for u in users:
        ws.append([u.get("id"), u.get("role"), u.get("fname"), u.get("lname"),
                   u.get("email"), u.get("contact_number"), u.get("company_name"),
                   u.get("payment_status"), str(u.get("created_at", ""))])
    audit_log(session.get("user_id"), "export_report", "report", None, "Downloaded users.xlsx")
    return _xl_response(wb, "users.xlsx")


@admin_bp.route("/admin/export/payments")
def export_payments():
    resp = require_admin()
    if resp:
        return resp
    from openpyxl import Workbook
    conn = get_db(); cur = dict_cur(conn)
    cur.execute("""
        SELECT p.id, u.fname, u.lname, u.email, p.amount_php, p.method,
               p.status, p.receipt_id, p.tx_hash, p.created_at, p.paid_at
        FROM payments p JOIN users u ON p.user_id = u.id ORDER BY p.created_at DESC
    """)
    payments = cur.fetchall(); cur.close(); conn.close()
    wb = Workbook(); ws = wb.active; ws.title = "Payments"
    headers = ["ID", "First Name", "Last Name", "Email", "Amount (PHP)", "Method",
               "Status", "Receipt ID", "TX Hash", "Created At", "Paid At"]
    _xl_header(ws, headers)
    for p in payments:
        ws.append([p["id"], p["fname"], p["lname"], p["email"],
                   float(p["amount_php"]) if p["amount_php"] else 0,
                   p["method"], p["status"], p.get("receipt_id"),
                   p.get("tx_hash"), str(p["created_at"]), str(p.get("paid_at") or "")])
    audit_log(session.get("user_id"), "export_report", "report", None, "Downloaded payments.xlsx")
    return _xl_response(wb, "payments.xlsx")


@admin_bp.route("/admin/export/lms-progress")
def export_lms_progress():
    resp = require_admin()
    if resp:
        return resp
    from openpyxl import Workbook
    conn = get_db(); cur = dict_cur(conn)
    cur.execute("""
        SELECT u.fname, u.lname, u.email, m.title AS module,
               p.quiz_score, p.quiz_passed, p.exam_grade,
               p.completed, p.completed_at
        FROM lms_progress p
        JOIN users u ON p.user_id = u.id
        JOIN lms_modules m ON p.module_id = m.id
        ORDER BY u.lname, m.title
    """)
    rows = cur.fetchall(); cur.close(); conn.close()
    wb = Workbook(); ws = wb.active; ws.title = "LMS Progress"
    headers = ["First Name", "Last Name", "Email", "Module",
               "Quiz Score", "Quiz Passed", "Exam Grade", "Completed", "Completed At"]
    _xl_header(ws, headers)
    for r in rows:
        ws.append([r["fname"], r["lname"], r["email"], r["module"],
                   float(r["quiz_score"]) if r["quiz_score"] else None,
                   "Yes" if r["quiz_passed"] else "No",
                   r["exam_grade"],
                   "Yes" if r["completed"] else "No",
                   str(r["completed_at"] or "")])
    audit_log(session.get("user_id"), "export_report", "report", None, "Downloaded lms-progress.xlsx")
    return _xl_response(wb, "lms-progress.xlsx")


@admin_bp.route("/admin/export/solar-requests")
def export_solar_requests():
    resp = require_admin()
    if resp:
        return resp
    from openpyxl import Workbook
    reqs = get_solar_requests()
    wb = Workbook(); ws = wb.active; ws.title = "Solar Requests"
    headers = ["ID", "Name", "Email", "Contact", "Establishment",
               "Monthly Bill (PHP)", "System Size (kW)", "Status", "Created At"]
    _xl_header(ws, headers)
    for r in reqs:
        ws.append([r.get("id"), r.get("name"), r.get("email"), r.get("contact"),
                   r.get("establishment_type"),
                   float(r.get("monthly_bill_php") or 0),
                   float(r.get("system_size_kw") or 0),
                   r.get("status"), str(r.get("created_at", ""))])
    audit_log(session.get("user_id"), "export_report", "report", None, "Downloaded solar-requests.xlsx")
    return _xl_response(wb, "solar-requests.xlsx")
